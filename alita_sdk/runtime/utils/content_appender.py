"""Handle binary document formats: append and create.

Provides registries of format-specific handlers and dispatchers
that the Artifact client (or any other caller) can use without knowing
format internals.
"""

import json
import logging
import re
from io import BytesIO
from pathlib import Path

logger = logging.getLogger(__name__)

# Simple regex to detect HTML content vs Markdown/plain text.
_HTML_TAG_RE = re.compile(
    r'<(?:html|body|div|table|h[1-6]|p|ul|ol|li|span|br|hr|img|a\s)'
    r'[\s>/]',
    re.IGNORECASE,
)


def _looks_like_html(text: str) -> bool:
    """Return True when *text* appears to contain HTML markup."""
    return bool(_HTML_TAG_RE.search(text))


# ---------------------------------------------------------------------------
# Format handlers - create and append operations
# ---------------------------------------------------------------------------

def _convert_and_append_html(text: str, doc) -> None:
    """Convert HTML/Markdown text and append rich content to a python-docx Document."""
    from htmldocx import HtmlToDocx

    html = text
    if not _looks_like_html(text):
        import markdown
        html = markdown.markdown(
            text, extensions=['tables', 'fenced_code']
        )

    parser = HtmlToDocx()
    parser.table_style = 'TableGrid'
    parser.add_html_to_document(html, doc)


def _create_docx(text: str) -> bytes:
    """Create a DOCX binary from HTML or Markdown/plain-text content."""
    from docx import Document as DocxDocument

    doc = DocxDocument()
    _convert_and_append_html(text, doc)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def _append_docx(raw_bytes: bytes, text: str) -> bytes:
    """Append rich content (HTML/Markdown/plain text) to a DOCX file."""
    from docx import Document as DocxDocument

    doc = DocxDocument(BytesIO(raw_bytes))
    _convert_and_append_html(text, doc)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


_OLE2_MAGIC = b'\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1'


def _append_doc(raw_bytes: bytes, text: str) -> bytes:
    """Append text inside HTML body tags for .doc files (HTML format only).

    Raises ValueError for binary OLE2 .doc files — convert to .docx first.
    """
    if raw_bytes[:8] == _OLE2_MAGIC:
        raise ValueError(
            "Cannot append to a binary OLE2 .doc file. "
            "Only HTML-based .doc files are supported. "
            "Convert the file to .docx first."
        )

    from bs4 import BeautifulSoup

    # Parse existing HTML
    soup = BeautifulSoup(raw_bytes.decode('utf-8', errors='ignore'), 'html.parser')
    
    # Find or create body tag
    body = soup.find('body')
    if not body:
        # Malformed HTML without body - wrap everything and append
        html_tag = soup.find('html')
        if html_tag:
            body = soup.new_tag('body')
            # Move all children of html into body
            for child in list(html_tag.children):
                body.append(child)
            html_tag.append(body)
        else:
            # No html/body structure at all - create minimal structure
            new_html = soup.new_tag('html')
            body = soup.new_tag('body')
            # Move all existing content into body
            for child in list(soup.children):
                body.append(child)
            new_html.append(body)
            soup.append(new_html)
    
    # Append new content as paragraphs inside body
    for line in text.split('\n'):
        if line.strip():  # Skip empty lines
            p = soup.new_tag('p')
            p.string = line
            body.append(p)
    
    return str(soup).encode('utf-8')


def _create_xlsx(text: str) -> bytes:
    """Create an Excel file from JSON data.
    
    Expects JSON string with structure:
    {
        "Sheet1": [["A1", "B1"], ["A2", "B2"]],
        "Sheet2": [["Data"]]
    }
    """
    from openpyxl import Workbook
    
    try:
        data = json.loads(text)
    except json.JSONDecodeError as e:
        raise ValueError(f"Invalid JSON format for .xlsx file data: {e}")
    
    if not isinstance(data, dict):
        raise ValueError(f"Excel data must be a dict with sheet names as keys, got {type(data).__name__}")
    
    try:
        workbook = Workbook()
        first_sheet = True
        
        for sheet_name, sheet_data in data.items():
            if first_sheet:
                sheet = workbook.active
                sheet.title = sheet_name
                first_sheet = False
            else:
                sheet = workbook.create_sheet(title=sheet_name)
            
            for row in sheet_data:
                sheet.append(row)
        
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
        
    except Exception as e:
        raise ValueError(f"Error processing .xlsx file data: {e}")


# ---------------------------------------------------------------------------
# Format registry - single source of truth for write operations
# ---------------------------------------------------------------------------
# Each format should provide create and/or append handlers based on its needs:
# - Both create + append: Full lifecycle support (e.g., .docx)
# - Only create: Write-once formats or formats where append logic not yet implemented
# - Only append: Formats already generated correctly by LLM (e.g., .doc = HTML)
#
# Structure: {extension: {
#     'create': handler(text) -> bytes | None (None = pass-through, LLM generates correct format),
#     'append': handler(raw_bytes, text) -> bytes | None (None = append not supported yet),
#     'tool_description_line': str (one-line description for create_file tool),
#     'param_description_block': str (detailed description for filedata parameter)
# }}

_format_registry = {
    '.xlsx': {
        'create': _create_xlsx,
        'append': None,  # Append to xlsx not yet implemented
        'tool_description_line': '- .xlsx: Excel spreadsheet with multiple sheets and tabular data. Provide filedata as JSON.',
        'param_description_block': '''
For .xlsx files, provide JSON with sheet names as keys and 2D arrays as values. Example:
{
    "Sheet1": [
        ["Name", "Age", "City"],
        ["Alice", 25, "New York"],
        ["Bob", 30, "San Francisco"]
    ],
    "Sheet2": [["Data"]]
}'''
    },
    '.docx': {
        'create': _create_docx,
        'append': _append_docx,
        'tool_description_line': '- .docx: Word document with tables, headings, formatting. Provide filedata as HTML or Markdown.',
        'param_description_block': '''
For .docx files, provide HTML or Markdown content:
- HTML example: "<h1>Report</h1><p>Text.</p><table><tr><th>Name</th><th>Score</th></tr><tr><td>Alice</td><td>95</td></tr></table>"
- Markdown example: "# Report\\n\\nText.\\n\\n| Name | Score |\\n|------|-------|\\n| Alice | 95 |"'''
    },
    '.doc': {
        'create': None,  # LLM already generates HTML - no conversion needed
        'append': _append_doc,
        'tool_description_line': None,  # No create handler, so no tool description
        'param_description_block': None  # No create handler, so no parameter description
    },
}


_creators = {ext: meta['create'] for ext, meta in _format_registry.items() if meta['create'] is not None}
_appenders = {ext: meta['append'] for ext, meta in _format_registry.items() if meta['append'] is not None}


def append_to_binary(filename: str, raw_bytes: bytes, text: str) -> bytes | None:
    """Try to append text to a binary file, preserving its format.

    Returns modified file bytes or None (caller falls back to text append).
    """
    ext = Path(filename).suffix.lower()
    handler = _appenders.get(ext)
    if handler is None:
        return None
    try:
        return handler(raw_bytes, text)
    except Exception as e:
        logger.error("Failed to append to %s file '%s': %s",
                     ext, filename, e)
        raise


def create_from_content(filename: str, text: str) -> bytes | None:
    """Try to create a binary file from text content.

    Returns file bytes or None (caller stores the text as-is).
    """
    ext = Path(filename).suffix.lower()
    handler = _creators.get(ext)
    if handler is None:
        return None
    try:
        return handler(text)
    except Exception as e:
        logger.error("Failed to create %s file '%s': %s",
                     ext, filename, e)
        raise


def get_tool_description_lines(handler_key: str = 'create') -> list[str]:
    """Return description lines for formats that have the given handler."""
    return [
        meta['tool_description_line']
        for meta in _format_registry.values()
        if meta.get(handler_key) is not None
        and meta.get('tool_description_line') is not None
    ]


def get_param_description_blocks(handler_key: str = 'create') -> list[str]:
    """Return description blocks for formats that have the given handler."""
    return [
        meta['param_description_block']
        for meta in _format_registry.values()
        if meta.get(handler_key) is not None
        and meta.get('param_description_block') is not None
    ]
